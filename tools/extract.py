# -*- coding: utf-8 -*-
"""Извлечение данных из Excel-книг в data/default_config.json.
Запуск: python tools/extract.py
Скрипт ОБЯЗАН падать (assert), если какой-либо эффект технологий / D60 не распознан."""
import glob
import json
import os
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "source", "tablets", "2026-10-07")
GM_BOOK = os.path.join(ROOT, "source", "gm-book", "Книга_игромастера_Глобальный_ИИ_v0.3_2026-10-07.xlsx")
OUT = os.path.join(ROOT, "data", "default_config.json")

# ---------- справочники ----------
COUNTRIES = [  # (id, имя, ISO numeric) в порядке листа «Распределение команд»
    ("russia", "Россия", 643), ("usa", "США", 840), ("china", "Китай", 156),
    ("india", "Индия", 356), ("uk", "Великобритания", 826), ("france", "Франция", 250),
    ("germany", "Германия", 276), ("japan", "Япония", 392), ("skorea", "Южная Корея", 410),
    ("turkey", "Турция", 792), ("saudi", "Саудовская Аравия", 682), ("iran", "Иран", 364),
    ("israel", "Израиль", 376), ("brazil", "Бразилия", 76), ("uae", "ОАЭ", 784),
    ("indonesia", "Индонезия", 360),
]
CID = {n: i for i, n, _ in COUNTRIES}

RES_IDS = [
    "minerals", "chem", "energy", "water", "space", "semis", "compute", "dc", "network", "devices",
    "sw_base", "sw_dev", "sw_ops", "math", "methods",
    "data_raw", "data_derived", "data_context", "model_res", "eval", "docs", "access", "cyber",
    "human", "expertise", "digitized", "integration", "org_arch", "org_change", "finance", "providers", "econ_viab",
    "law", "regulatory", "institutions", "trade_env", "trust", "workforce", "local_fit", "ai_gov",
    "monitoring", "renewal", "recycling",
]
SPH_IDS = ["production", "management", "cognition", "learning", "war", "economy", "gov", "adoption"]
GROUP_ROWS = {"I": (8, 17), "II": (19, 23), "III": (25, 32), "IV": (34, 42), "V": (44, 51), "VI": (53, 55)}

# псевдонимы: (основа без окончания, цель). Сопоставление по началу фрагмента, побеждает самая длинная основа.
ALIASES = [
    ("минерально-сырьев", "res:minerals"), ("химическ", "res:chem"),
    ("энергетическ", "res:energy"), ("энергетик", "res:energy"),
    ("вода", "res:water"), ("воды", "res:water"),
    ("пространственн", "res:space"), ("полупроводник", "res:semis"),
    ("мощност", "res:semis"),
    ("вычислительн", "res:compute"), ("цод", "res:dc"),
    ("сетев", "res:network"), ("сети", "res:network"), ("конечн", "res:devices"),
    ("базов", "res:sw_base"), ("стек разработки", "res:sw_dev"), ("стека разработки", "res:sw_dev"),
    ("стек эксплуатации", "res:sw_ops"), ("стека эксплуатации", "res:sw_ops"),
    ("математико", "res:math"), ("метод", "res:methods"),
    ("исходн", "res:data_raw"), ("производн", "res:data_derived"), ("контекстн", "res:data_context"),
    ("модельн", "res:model_res"),
    ("ресурсы оценки", "res:eval"), ("ресурсов оценки", "res:eval"), ("ресурс оценки", "res:eval"),
    ("документац", "res:docs"), ("права доступа", "res:access"), ("прав доступа", "res:access"),
    ("кибербезопасн", "res:cyber"),
    ("человеческ", "res:human"), ("квалифицированн", "res:human"),
    ("предметн", "res:expertise"), ("оцифрован", "res:digitized"), ("интеграционн", "res:integration"),
    ("организационная архитектур", "res:org_arch"), ("организационной архитектур", "res:org_arch"),
    ("организационная способност", "res:org_change"), ("организационной способност", "res:org_change"),
    ("финансов", "res:finance"), ("поставщик", "res:providers"),
    ("экономическая целесообразн", "res:econ_viab"), ("экономической целесообразн", "res:econ_viab"),
    ("правов", "res:law"), ("регуляторн", "res:regulatory"), ("институциональн", "res:institutions"),
    ("конкурентн", "res:trade_env"), ("торговая среда", "res:trade_env"), ("торговой среды", "res:trade_env"),
    ("социальное доверие", "res:trust"), ("социального доверия", "res:trust"),
    ("готовность работников", "res:workforce"), ("готовности работников", "res:workforce"),
    ("локальн", "res:local_fit"),
    ("система управления ии", "res:ai_gov"), ("системы управления ии", "res:ai_gov"),
    ("поствнедренческ", "res:monitoring"), ("ресурс обновления", "res:renewal"),
    ("ресурс утилизации", "res:recycling"),
    # сферы
    ("производств", "sph:production"), ("управлени", "sph:management"), ("познани", "sph:cognition"),
    ("обучени", "sph:learning"), ("войн", "sph:war"), ("экономик", "sph:economy"),
    ("госуправлен", "sph:gov"), ("государственн", "sph:gov"), ("принят", "sph:adoption"),
]
ALIASES.sort(key=lambda a: -len(a[0]))
GOD = {"арес": "ares", "афин": "athena", "аполлон": "apollo"}
SIGN = {"−": "-", "–": "-", "+": "+", "-": "-"}


def norm(s):
    return s.lower().replace("ё", "е").strip()


def num(s):
    s = s.replace("−", "-").replace("–", "-").replace(",", ".")
    return float(s)


def find_target(text):
    t = norm(text)
    for stem, tgt in ALIASES:
        if t.startswith(stem):
            return tgt
    return None


def gods_in(t):
    return [v for k, v in GOD.items() if k in t]


# ---------- парсер модификаторов (проценты) ----------
def parse_mod(part):
    """Возвращает список {target:'mod:..',kind:'MOD',value} или None, если это не модификатор."""
    t = norm(part)
    m = re.search(r"([+\-−]?)\s*(\d+(?:[.,]\d+)?)\s*%", t)
    if "открывает карты" in t:
        return [{"target": "mod:x4_double_swan", "kind": "MOD", "value": 1}]
    if not m:
        return None
    v = num((m.group(1) or "+") + m.group(2)) / 100.0
    if "вся оборона" in t:
        keys = ["all_def"]
    elif "очки трансформации" in t and "групп" in t:
        keys = ["transform_group_I"]
    elif "очки ии в сферы" in t:
        keys = ["ai_points_spheres"]
    elif "эффективность очков ии" in t:
        keys = ["ai_points_total"]
    elif "ветви модели" in t:
        keys = ["science_branch_M"]
    elif "науки для m/a/x" in t:
        keys = ["science_branch_MAX"]
    elif "очков науки" in t or "эффективность науки" in t:
        keys = ["science_all"]
    elif "улучшение отношений" in t:
        keys = ["relations_improve"]
    elif t.startswith("защит") or "защита" in t:
        keys = [g + "_def" for g in gods_in(t)]
    elif t.startswith("атак"):
        keys = [g + "_att" for g in gods_in(t)]
    elif t.startswith("эффективность защит"):
        keys = [g + "_def" for g in gods_in(t)]
    elif t.startswith("эффективность"):
        keys = [g + s for g in gods_in(t) for s in ("_att", "_def")]
    else:
        raise ValueError("не распознан модификатор: " + part)
    assert keys, "нет ключей для модификатора: " + part
    return [{"target": "mod:" + k, "kind": "MOD", "value": round(v, 4)} for k in keys]


# ---------- парсер технологий ----------
def parse_tech_effects(text, overrides=None):
    parts = [p.strip() for p in text.split(";") if p.strip()]
    effects, run = [], []  # run: цели последних ПЗ подряд
    last = None
    for p in parts:
        mods = parse_mod(p)
        if mods is not None:
            effects += mods
            continue
        t = norm(p)
        m = re.match(r"^(.*?)\s*\+\s*(\d+(?:[.,]\d+)?)(/год)?$", t)
        assert m, "не удалось разобрать часть: " + p
        name, val, per = m.group(1), num(m.group(2)), bool(m.group(3))
        pron = re.match(r"^(его|её|ее|их)\s+прирост$", name)
        if pron:
            assert per and last is not None, "прирост без цели/в год: " + p
            targets = run if pron.group(1) == "их" else [last]
            for tg in targets:
                effects.append({"target": tg, "kind": "PP", "value": val})
            continue
        if name.startswith("прирост "):
            tg = find_target(name[len("прирост "):])
            assert tg, "цель не найдена: " + p
            effects.append({"target": tg, "kind": "PP", "value": val})
            last = tg
            run = [tg]
            continue
        tg = find_target(name)
        assert tg, "цель не найдена: " + p
        if per:  # «Производных +0,15/год» — продолжение списка приростов
            effects.append({"target": tg, "kind": "PP", "value": val})
            last = tg
            run = [tg]
        else:
            prev_pz = bool(effects) and effects[-1]["kind"] == "PZ"
            effects.append({"target": tg, "kind": "PZ", "value": val})
            last = tg
            run = (run if prev_pz else []) + [tg]
    return effects


# ---------- парсер D60 ----------
KINDS = {"ПЗ": "PZ", "ВЗ": "VZ", "ПП": "PP", "ВП": "VP"}


def parse_swan_effects(text):
    t = text.strip().rstrip(".")
    t = re.sub(r"^(Все играющие страны|Каждой затронутой стране):\s*", "", t)
    parts = [p.strip() for p in t.split(";") if p.strip()]
    effects, last = [], None
    for p in parts:
        mods = parse_mod(p)
        if mods is not None:
            for md in mods:  # D60: на 1 ход
                md["ttl"] = 1
            effects += mods
            continue
        m = re.match(r"^всем 8 сферам(?: внедрения ИИ)? (ВЗ|ПЗ|ВП|ПП)\s*([+\-−]\d+(?:,\d+)?)$", p)
        if m:
            effects.append({"target": "all_spheres", "kind": KINDS[m.group(1)], "value": num(m.group(2))})
            continue
        m = re.match(r"^(?:(.+?):\s*)?(ПЗ|ВЗ|ПП|ВП)\s*([+\-−]\d+(?:,\d+)?)$", p)
        assert m, "D60: не удалось разобрать часть: " + p
        name, kind, val = m.group(1), m.group(2), num(m.group(3))
        if name is None or re.match(r"^(его|её|их)\s+прирост$", norm(name)):
            assert last, "D60: нет предыдущей цели: " + p
            tg = last
        else:
            tg = find_target(name)
            assert tg, "D60: цель не найдена: " + p
        effects.append({"target": tg, "kind": KINDS[kind], "value": val})
        last = tg
    return effects


# ---------- чтение книг ----------
def find_sheet(wb):
    for s in wb.sheetnames:
        n = s.strip().lower()
        if n.endswith(" 2026") and not n.startswith(("ход", "методика")):
            return wb[s]
    raise RuntimeError("нет листа «<Страна> 2026»")


def read_country(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = find_sheet(wb)
    res, sph, rel = {}, {}, {}
    names = []
    for gi, (g, (a, b)) in enumerate(GROUP_ROWS.items()):
        pass
    rows = [r for g, (a, b) in GROUP_ROWS.items() for r in range(a, b + 1)]
    assert len(rows) == 43
    for rid, r in zip(RES_IDS, rows):
        B, C = ws.cell(r, 2).value, ws.cell(r, 3).value
        assert isinstance(B, (int, float)) and isinstance(C, (int, float)), (path, r, B, C)
        res[rid] = {"B": float(B), "C": float(C)}
        names.append(ws.cell(r, 1).value)
    for sid, r in zip(SPH_IDS, range(11, 19)):
        L, M = ws.cell(r, 12).value, ws.cell(r, 13).value
        assert isinstance(L, (int, float)) and isinstance(M, (int, float)), (path, r)
        sph[sid] = {"L": float(L), "M": float(M)}
    sph_names = [ws.cell(r, 11).value for r in range(11, 19)]
    for r in range(11, 27):
        n, v = ws.cell(r, 21).value, ws.cell(r, 22).value
        if n == "Прочее":
            continue
        assert n in CID, (path, n)
        rel[CID[n]] = float(v)
    assert len(rel) == 15
    return wb, ws, res, sph, rel, names, sph_names


def main():
    files = glob.glob(os.path.join(SRC, "Игра_ИИ_*"))
    books = {}
    for cid, name, iso in COUNTRIES:
        fn = name.replace(" ", "_")
        cand = [f for f in files if os.path.basename(f).startswith(f"Игра_ИИ_{fn}_2026")]
        assert len(cand) == 1, (name, cand)
        books[cid] = cand[0]

    cfg = {"version": 1}
    countries, ref_names, ref_sph, ref_tech = [], None, None, None
    usa_wb = None
    for cid, name, iso in COUNTRIES:
        wb, ws, res, sph, rel, names, sph_names = read_country(books[cid])
        if ref_names is None:
            ref_names, ref_sph = names, sph_names
        assert names == ref_names and sph_names == ref_sph, "структура книги " + name + " отличается"
        t = wb["Технологии"]
        tech_sig = [tuple(t.cell(r, c).value for c in (1, 2, 3, 4, 5, 6, 7, 8, 9, 10)) for r in range(11, 51)]
        if ref_tech is None:
            ref_tech = tech_sig
        assert tech_sig == ref_tech, "технологии " + name + " отличаются"
        assert set(rel) == {c for c, _, _ in COUNTRIES} - {cid}, name
        countries.append({"id": cid, "name": name, "isoNumeric": iso,
                          "start": {"resources": res, "spheres": sph}, "relations": rel})
        if cid == "usa":
            usa_wb = wb
    print("Структура 16 книг идентична: ресурсы, сферы, технологии. OK")

    groups = {}
    for g, (a, b) in GROUP_ROWS.items():
        for _ in range(a, b + 1):
            pass
    flat_groups = [g for g, (a, b) in GROUP_ROWS.items() for _ in range(a, b + 1)]
    cfg["resources"] = [{"id": i, "name": n, "group": g} for i, n, g in zip(RES_IDS, ref_names, flat_groups)]
    cfg["spheres"] = [{"id": i, "name": n} for i, n in zip(SPH_IDS, ref_sph)]
    cfg["countries"] = countries
    assert len(cfg["resources"]) == 43 and len(cfg["spheres"]) == 8

    # ---- технологии ----
    techs, fails = [], []
    for row in ref_tech:
        tid, branch, level, tname, cost, p1, p2, p3, essence, eff_text = row
        try:
            eff = parse_tech_effects(eff_text)
        except Exception as e:  # noqa
            fails.append((tid, eff_text, str(e)))
            eff = []
        techs.append({"id": tid, "branch": branch, "level": level, "name": tname, "cost": float(cost),
                      "prereqs": [p for p in (p1, p2, p3) if p], "essence": essence,
                      "effectText": eff_text, "effects": eff})
    cfg["techs"] = techs

    # ---- D60 ----
    gm = openpyxl.load_workbook(GM_BOOK, data_only=True)
    ws = gm["Случайные события D60"]
    tmap = {"Общемировое": "global", "Для бросившей страны": "national", "Страна + реальные союзники": "bloc"}
    swans = []
    for r in range(2, 62):
        n, typ, sign, title, text = [ws.cell(r, c).value for c in range(27, 32)]
        assert n == r - 1 and typ in tmap, (r, n, typ)
        try:
            eff = parse_swan_effects(text)
        except Exception as e:  # noqa
            fails.append((f"D60#{n}", text, str(e)))
            eff = []
        swans.append({"id": n, "type": tmap[typ], "sign": "+" if sign == "+" else "-", "title": title,
                      "text": text, "effects": eff})
    cfg["blackSwans"] = swans

    # ---- ручные override-ы (если автопарсинг что-то не осилил) ----
    OVERRIDES = {}  # id -> список эффектов
    for tid, text, err in list(fails):
        if tid in OVERRIDES:
            for coll in (techs, swans):
                for o in coll:
                    if str(o["id"]) == tid.replace("D60#", ""):
                        o["effects"] = OVERRIDES[tid]
            fails.remove((tid, text, err))
    if fails:
        print("НЕ РАСПАРСИЛОСЬ:")
        for f in fails:
            print("  ", f)
        raise SystemExit(1)

    # ---- проверка: каждый эффект → известная цель, каждая часть текста дала эффект ----
    rid_set, sid_set = set(RES_IDS), set(SPH_IDS)
    mod_keys = {"ares_att", "athena_att", "apollo_att", "ares_def", "athena_def", "apollo_def", "all_def",
                "relations_improve", "science_all", "science_branch_M", "science_branch_MAX", "ai_points_spheres",
                "transform_group_I", "ai_points_total", "x4_double_swan"}

    def check(owner, text, effs):
        assert effs, f"{owner}: пустые эффекты ({text})"
        for e in effs:
            tg = e["target"]
            ok = (tg.startswith("res:") and tg[4:] in rid_set) or (tg.startswith("sph:") and tg[4:] in sid_set) \
                or tg == "all_spheres" or (tg.startswith("mod:") and tg[4:] in mod_keys)
            assert ok, f"{owner}: неизвестная цель {tg} ({text})"
            assert e["kind"] in ("PZ", "VZ", "PP", "VP", "MOD"), (owner, e)
            assert isinstance(e["value"], (int, float))
        parts = [p for p in re.split(r";", re.sub(r"^(Все играющие страны|Каждой затронутой стране):\s*", "", text)) if p.strip()]
        # частей может быть меньше, чем эффектов («защита Афины и Аполлона»), но не больше
        assert len(effs) >= len(parts), f"{owner}: потеряны части текста ({text})"

    for t in techs:
        check(t["id"], t["effectText"], t["effects"])
    for s in swans:
        check(f"D60#{s['id']}", s["text"], s["effects"])
    print(f"Эффекты: {len(techs)} технологий и {len(swans)} событий D60 распознаны полностью.")
    for t in techs:
        print("  ", t["id"], t["effectText"], "=>", [f'{e["target"]}:{e["kind"]}:{e["value"]}' for e in t["effects"]])
    for s in swans:
        print("  D60", s["id"], "=>", [f'{e["target"]}:{e["kind"]}:{e["value"]}' for e in s["effects"]])

    # ---- лестница статусов из формулы W11 ----
    wbf = openpyxl.load_workbook(books["usa"])
    f = wbf[find_sheet(wbf).title]["W11"].value
    ladder = []
    for m in re.finditer(r'IF\(V11(<=|<|=)(-?[\d.]+),"([^"]+)"', f):
        op, thr, label = m.groups()
        key = {"<=": "le", "<": "lt", "=": "eq"}[op]
        ladder.append({key: float(thr), "label": label})
    ladder.append({"else": True, "label": re.findall(r'"([^"]+)"\)+$', f)[0]})
    assert len(ladder) == 11, ladder
    cfg["ladder"] = ladder

    # ---- формулы (JS-выражения; контекст: res, sph, rel, pow, ln, min, max, ces, geo, sign, abs, exp) ----
    R = lambda *ids: "[" + ",".join(f"'{i}'" for i in ids) + "]"
    cfg["formulas"] = {
        "_help": "Тела JS-функций. Контекст: res{id}, sph{id} (реальные значения), rel[15 кодов], pow, ln, min, max, "
                 "sign, abs, exp, ces(ids,rho) — CES-среднее ресурсов (значение/10), geo({id:вес}) — взвешенное "
                 "геометрическое по ресурсам и сферам (значение/10).",
        "ai": {
            "x": "const g1=ces(" + R(*RES_IDS[0:10]) + ",-1), g2=ces(" + R(*RES_IDS[10:15]) + ",-0.5), g3=ces(" +
                 R(*RES_IDS[15:23]) + ",-0.5), g4=ces(" + R(*RES_IDS[23:32]) + ",-0.25), g5=ces(" +
                 R(*RES_IDS[32:40]) + ",0.25), g6=geo({monitoring:1/3,renewal:1/3,recycling:1/3});\n"
                 "return pow(g1,0.3)*pow(g2,0.2)*pow(g3,0.2)*pow(g4,0.15)*pow(g5,0.1)*pow(g6,0.05);",
            "n": 3, "k": 0.55, "normalize": True},
        "science": {
            "x": "return pow(geo({cognition:0.45,learning:0.25,management:0.1,gov:0.1,economy:0.05,adoption:0.05}),0.8)"
                 "*pow(geo({human:0.2,math:0.15,data_raw:0.1,data_context:0.1,dc:0.15,model_res:0.1,finance:0.1,access:0.1}),0.2);",
            "n": 2, "k": 0.55, "normalize": True},
        "fp": "return 100*pow((rel.reduce((a,b)=>a+b,0)/rel.length+5)/10,0.3)"
              "*pow(geo({network:0.15,cyber:0.15,finance:0.2,providers:0.15,institutions:0.15,trade_env:0.2}),0.3)"
              "*pow(geo({management:0.2,war:0.2,economy:0.3,gov:0.3}),0.4);",
        "transform": "const s=pow(geo({production:0.2,management:0.15,cognition:0.1,learning:0.1,war:0.05,economy:0.15,gov:0.15,adoption:0.1}),0.65);\n"
                     "const r=pow(geo({minerals:0.06,chem:0.1,energy:0.12,water:0.05,space:0.04,semis:0.17,compute:0.15,dc:0.15,network:0.1,devices:0.06}),0.35);\n"
                     "return 100*ln(1+20*s*r)/ln(21);",
        "growth": {
            "_help": "Переход года. Переменные: B,C,L,M — базовые значения/приросты; I,S — реальные значения; "
                     "b,c,h — вложенные очки; vp — ВП; dPZ,dPP — новые постоянные эффекты; kI,kAI — модификаторы.",
            "resB": "B + (C + vp) + 0.35*ln(1+0.6*b*kI)*pow(max(0,1-I/10),0.7) + dPZ",
            "resC": "min(1, C + 0.08*ln(1+0.5*c)*max(0,1-C) + dPP)",
            "sphL": "L + (M + vp) + dPZ",
            "sphM": "min(1, M + 0.1*ln(1+0.5*h*kAI)*pow(max(0,1-S/10),0.8) + dPP)",
        },
    }
    cfg["diplomacy"] = {
        "_help": "Шанс NPC. Переменные: M — мой код отношений, w — эффект вложений, sg — sign(M)*(|M|/5)^1.25.",
        "chanceImprove": "M>=5 ? 0 : min(0.75, 0.1+0.08*sg+0.5*w*((5-M)/10))",
        "chanceWorsen": "M<=-5 ? 0 : max(0, (0.1-0.08*sg)*(1-0.85*w))",
        "wDenominator": 40,
        "improveCost": 5,          # цена +1 к отношениям с игроком (ВП-очки)
        "allianceMin": 1,          # союз с игроком возможен, если код > allianceMin
        "npcAllianceOn": 3,        # союз с NPC создаётся при коде > 3
        "npcAllianceOff": 2,       # теряется при коде <= 2
        "npcExchange": {"rate": 2, "minRelation": 0, "maxReceivePerNpcPerTurn": 1.0},
        "minCode": -5, "maxCode": 5,
    }
    cfg["war"] = {
        "power": {"ares": 1.0, "athena": 0.5, "apollo": 0.25},
        "drawFactor": 0.8, "dieSize": 60,
        "totalLevels": {"resource": [[0.20, 0.01], [0.30, 0.02], [0.40, 0.03]],
                        "sphere": [0.02, 0.03, 0.05], "relation": [0.2, 0.3, 0.5]},
        "local": {"minDiff": 5, "ratios": [2, 3, 5],
                  "ares": {"value": [0.15, 0.20, 0.30], "growth": [0.01, 0.02, 0.03], "group": "I"},
                  "athena": [0.02, 0.03, 0.05], "apollo": [0.2, 0.3, 0.5]},
        "npcDefense": {"ares": 8, "athena": 8, "apollo": 8},
        "relationDropOnAttack": 1,
        # v0.3: войны только против игроков; True возвращает поведение v0.2
        "allowNpcTargets": False,
    }
    cfg["timers"] = {"turnSeconds": 600, "draftSeconds": 120, "startYear": 2026, "endYear": 2040, "offlineEndSeconds": 30}
    cfg["victory"] = {
        "scales": {"enabled": True, "min": 100},
        "tech": {"enabled": True},
        "diplomacy": {"enabled": True},
        "tieBreak": "lowerStartSum",
    }
    cfg["limits"] = {"resMin": 1, "resMax": 10, "growthMax": 1, "maxPlayers": 16,
                     "diffusion": {"perCountry": 0.1, "minFactor": 0.5},
                     "tradeMaxAmount": 5, "tradeMaxYears": 5,
                     # v0.3: страна выдаётся случайно; "pick" возвращает драфт v0.2
                     "draftMode": "random"}

    # ---- эталон из Excel для selfcheck (США) ----
    wbv = usa_wb
    s26, s27 = wbv["США 2026"], wbv["США 2027"]
    rows = [r for g, (a, b) in GROUP_ROWS.items() for r in range(a, b + 1)]
    cfg["selfcheck"] = {
        "usa2026": {"ai": s26["K6"].value, "science": s26["N6"].value, "fp": s26["Q6"].value,
                    "transform": s26["U6"].value},
        "usa2027": {"res": {i: {"B": s27.cell(r, 2).value, "C": s27.cell(r, 3).value} for i, r in zip(RES_IDS, rows)},
                    "sph": {i: {"L": s27.cell(r, 12).value, "M": s27.cell(r, 13).value} for i, r in zip(SPH_IDS, range(11, 19))}},
    }
    print("Excel США 2026:", cfg["selfcheck"]["usa2026"])

    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(cfg, fh, ensure_ascii=False, indent=1)
    print("Записано:", OUT, os.path.getsize(OUT), "байт")


if __name__ == "__main__":
    main()
